"""前年実績(F-01 横展開) 仕掛品インポートサービス —
20 SC仕掛品.xlsx 「仕掛品SC明細」シートから 第38期の仕掛品(製造課)別 SC 原価フローを
prior_year_wip_actuals テーブルに upsert する。

シート構造 (仕掛品SC明細):
  R1   タイトル "第38期 仕掛品(製造課)"
  R3   ヘッダ (年度/番手/種類/名寄/SC単価/[数量|金額]/期首棚卸/...)
  R5〜 データ行 — 1仕掛品が「数量」行と「金額」行の2行ペア
       (年度/番手/種類/名寄/SC単価 は数量行のみに記載)
  末尾 合計行 + チェック行 (種類が空 or "数量チェック" 等) はスキップ。

列マップ (1-origin):
  1: 年度 (仕込年度 S62/H1 等)
  2: 番手
  3: 種類 (=wip_code, 一意キー 例 13R)
  4: 名寄 (原液グループ R/G 等)
  5: SC単価
  6: 行種別 (数量 / 金額)
  7:  期首棚卸
  8:  原材料
  9:  労務費
  10: 経費
  11: 前工程費
  12: 完成品
  13: 研究費
  14: 販促費
  15: 廃棄処分
  16: 次工程へ
  17: 製造部生産分
  18: 在庫調整
  19: 期末棚卸
  (c21以降は別表「原液単位原価」のため取込対象外)
"""

import io
import uuid
from datetime import datetime
from decimal import Decimal

from openpyxl import load_workbook
from sqlalchemy import delete
from sqlalchemy.ext.asyncio import AsyncSession

from app.models.audit import ImportBatch, ImportError as ImportErrorModel, ImportStatus
from app.models.cost import PriorYearWipActual, SourceSystem


SHEET_WIP = "仕掛品SC明細"

# 主要6項目: (prefix, 0-origin 列index)
MAIN_GROUPS: list[tuple[str, int]] = [
    ("opening", 6),
    ("material", 7),
    ("labor", 8),
    ("expense", 9),
    ("pre_process", 10),
    ("closing", 18),
]

# その他7項目: (項目名, JSONキー, 0-origin 列index)
OTHER_ITEM_COLS: list[tuple[str, str, int]] = [
    ("完成品", "finished", 11),
    ("研究費", "rnd", 12),
    ("販促費", "promo", 13),
    ("廃棄処分", "disposal", 14),
    ("次工程へ", "to_next", 15),
    ("製造部生産分", "mfg_dept", 16),
    ("在庫調整", "inventory_adj", 17),
]


def _to_decimal(v) -> Decimal:
    if v is None:
        return Decimal("0")
    if isinstance(v, str):
        s = v.strip()
        if not s or s in ("ー", "−", "-", "#REF!", "#N/A", "#VALUE!"):
            return Decimal("0")
        try:
            return Decimal(s)
        except Exception:
            return Decimal("0")
    if isinstance(v, bool):
        return Decimal("0")
    try:
        return Decimal(str(v))
    except Exception:
        return Decimal("0")


def _str_or_none(v) -> str | None:
    if v is None:
        return None
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    s = str(v).strip()
    return s if s else None


def parse_wip_sheet(content: bytes, sheet_name: str = SHEET_WIP) -> list[dict]:
    """仕掛品SC明細 シートをパースし、仕掛品行リストを返す。

    1仕掛品 = 「数量」行 + 「金額」行。数量行から qty を、金額行から cost を読む。
    種類(wip_code) が空の行 (合計行) や チェック行はスキップ。
    """
    wb = load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    ws = wb[sheet_name]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()

    out: list[dict] = []
    i = 4  # R5 (0-origin index 4)
    while i < len(rows):
        row = rows[i]
        if len(row) < 19:
            i += 1
            continue
        row_type = row[5]
        if row_type is None or str(row_type).strip() != "数量":
            i += 1
            continue
        wip_code = _str_or_none(row[2])  # 種類
        if not wip_code:  # 合計行 (種類が空)
            i += 1
            continue

        qty_row = row
        amount_row = rows[i + 1] if i + 1 < len(rows) else ()
        amt_type = amount_row[5] if len(amount_row) > 5 else None
        if amt_type is None or str(amt_type).strip() != "金額":
            amount_row = ()
            advance = 1
        else:
            advance = 2

        cost_items: dict[str, dict[str, str]] = {}
        for label, json_key, idx in OTHER_ITEM_COLS:
            qv = _to_decimal(qty_row[idx]) if len(qty_row) > idx else Decimal("0")
            cv = _to_decimal(amount_row[idx]) if len(amount_row) > idx else Decimal("0")
            cost_items[json_key] = {"label": label, "qty": str(qv), "cost": str(cv)}

        parsed = {
            "wip_code": wip_code,
            "production_year": _str_or_none(row[0]),
            "batch_no": _str_or_none(row[1]),
            "nayose": _str_or_none(row[3]),
            "sc_unit_price": _to_decimal(row[4]),
            "cost_items": cost_items,
        }
        for prefix, idx in MAIN_GROUPS:
            parsed[f"{prefix}_qty"] = _to_decimal(qty_row[idx]) if len(qty_row) > idx else Decimal("0")
            parsed[f"{prefix}_cost"] = _to_decimal(amount_row[idx]) if len(amount_row) > idx else Decimal("0")

        out.append(parsed)
        i += advance

    return out


async def process_prior_year_wip_import(
    db: AsyncSession,
    file_content: bytes,
    filename: str,
    fiscal_year: int,
    sheet_name: str = SHEET_WIP,
    source_system: str = "manual",
    delete_existing: bool = True,
    period_id: uuid.UUID | None = None,
) -> ImportBatch:
    """前年実績 仕掛品 Excel を取り込み、ImportBatch を返す。"""
    batch = ImportBatch(
        file_name=filename,
        source_system=source_system if source_system in SourceSystem._value2member_map_
        else SourceSystem.manual.value,
        status=ImportStatus.processing,
        period_id=period_id,
        total_rows=0,
        success_rows=0,
        error_rows=0,
        started_at=datetime.now(),
    )
    db.add(batch)
    await db.flush()

    try:
        rows = parse_wip_sheet(file_content, sheet_name)
    except Exception as e:
        batch.status = ImportStatus.failed
        batch.completed_at = datetime.now()
        batch.notes = f"ファイルパースエラー: {e}"
        db.add(ImportErrorModel(batch_id=batch.id, row_number=0, error_message=batch.notes))
        await db.flush()
        await db.refresh(batch)
        return batch

    batch.total_rows = len(rows)

    if delete_existing:
        await db.execute(
            delete(PriorYearWipActual).where(PriorYearWipActual.fiscal_year == fiscal_year)
        )
        await db.flush()

    success = 0
    seen: set[str] = set()
    dup = 0
    material_sum = Decimal("0")
    labor_sum = Decimal("0")
    expense_sum = Decimal("0")
    pre_sum = Decimal("0")
    close_sum = Decimal("0")

    for r in rows:
        if r["wip_code"] in seen:  # 同一 wip_code の重複は安全のためスキップ
            dup += 1
            continue
        seen.add(r["wip_code"])
        rec = PriorYearWipActual(
            fiscal_year=fiscal_year,
            wip_code=r["wip_code"],
            production_year=r["production_year"],
            batch_no=r["batch_no"],
            nayose=r["nayose"],
            sc_unit_price=r["sc_unit_price"],
            opening_qty=r["opening_qty"],
            opening_cost=r["opening_cost"],
            material_qty=r["material_qty"],
            material_cost=r["material_cost"],
            labor_qty=r["labor_qty"],
            labor_cost=r["labor_cost"],
            expense_qty=r["expense_qty"],
            expense_cost=r["expense_cost"],
            pre_process_qty=r["pre_process_qty"],
            pre_process_cost=r["pre_process_cost"],
            closing_qty=r["closing_qty"],
            closing_cost=r["closing_cost"],
            cost_items=r["cost_items"],
            source_file=filename,
            source_sheet=sheet_name,
            import_batch_id=batch.id,
        )
        db.add(rec)
        success += 1
        material_sum += r["material_cost"]
        labor_sum += r["labor_cost"]
        expense_sum += r["expense_cost"]
        pre_sum += r["pre_process_cost"]
        close_sum += r["closing_cost"]

    await db.flush()

    batch.success_rows = success
    batch.error_rows = 0
    batch.status = ImportStatus.completed
    batch.completed_at = datetime.now()
    batch.notes = (
        f"fiscal_year={fiscal_year}, total={len(rows)}, success={success}, dup_skipped={dup}; "
        f"material_sum={material_sum:.0f}, labor_sum={labor_sum:.0f}, "
        f"expense_sum={expense_sum:.0f}, pre_process_sum={pre_sum:.0f}, "
        f"closing_sum={close_sum:.0f}"
    )
    await db.flush()
    await db.refresh(batch)
    return batch
