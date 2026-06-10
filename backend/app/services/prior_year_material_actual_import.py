"""前年実績(F-01 横展開) 原材料インポートサービス —
10 SC原材料.xlsx 「原材料SC明細」シートから 第38期の原材料別 SC 原価フローを
prior_year_material_actuals テーブルに upsert する。

シート構造 (原材料SC明細):
  R1   タイトル "第38期 原材料"
  R3   ヘッダ (ｺｰﾄﾞ/原料名/SC単価/単位/期首棚卸/...)
  R5〜 データ行 — 1原料が「数量(単位)」行と「金額(円)」行の2行ペア
       (コード/原料名/SC単価/単位 は数量行のみに記載)

列マップ (1-origin):
  1: コード
  2: 原料名
  3: SC単価
  4: 単位 / 行種別ラベル (数量(㎏) or 金額(円))
  5:  期首棚卸
  6:  当期仕入高
  7:  当期投入高
  8:  検査・分析
  9:  試用
  10: 返品
  11: ロス分
  12: 廃棄処分
  13: 在庫調整
  14: その他調整
  15: 期末棚卸

末尾の「合計」行はスキップ。
"""

import io
import uuid
from datetime import datetime
from decimal import Decimal

from openpyxl import load_workbook
from sqlalchemy import delete, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.models.audit import ImportBatch, ImportError as ImportErrorModel, ImportStatus
from app.models.cost import PriorYearMaterialActual, SourceSystem
from app.models.master import Material


SHEET_MATERIAL = "原材料SC明細"

# 主要4項目: (JSON/列名 prefix, 0-origin 列index)
MAIN_GROUPS: list[tuple[str, int]] = [
    ("opening", 4),
    ("purchase", 5),
    ("input", 6),
    ("closing", 14),
]

# 調整系7項目: (Excel上の項目名, JSON キー, 0-origin 列index)
ADJUSTMENT_ITEM_COLS: list[tuple[str, str, int]] = [
    ("検査・分析", "inspection", 7),
    ("試用", "trial", 8),
    ("返品", "returns", 9),
    ("ロス分", "loss", 10),
    ("廃棄処分", "disposal", 11),
    ("在庫調整", "inventory_adj", 12),
    ("その他調整", "other_adj", 13),
]


def _to_decimal(v) -> Decimal:
    """セル値を Decimal に変換。空/エラー値は 0 とする。"""
    if v is None:
        return Decimal("0")
    if isinstance(v, str):
        s = v.strip()
        if not s or s in ("ー", "−", "-", "#REF!", "#N/A", "#VALUE!"):
            return Decimal("0")
        try:
            return Decimal(s)
        except Exception:
            return Decimal("0")
    try:
        return Decimal(str(v))
    except Exception:
        return Decimal("0")


def _normalize_code(v) -> str | None:
    """原料コードを文字列化。float表現 (1.0) は整数文字列に正規化。"""
    if v is None:
        return None
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    s = str(v).strip()
    return s if s else None


def _merge_row(existing: dict, new: dict) -> None:
    """同コードの2レコードを合算する (通常は発生しない安全策)。"""
    for prefix, _ in MAIN_GROUPS:
        existing[f"{prefix}_qty"] += new[f"{prefix}_qty"]
        existing[f"{prefix}_cost"] += new[f"{prefix}_cost"]
    for k, v in new["adjustment_items"].items():
        ex_v = existing["adjustment_items"].get(k)
        if ex_v is None:
            existing["adjustment_items"][k] = v
            continue
        existing["adjustment_items"][k] = {
            "label": ex_v.get("label") or v.get("label"),
            "qty": str(Decimal(ex_v["qty"]) + Decimal(v["qty"])),
            "cost": str(Decimal(ex_v["cost"]) + Decimal(v["cost"])),
        }


def parse_material_sheet(content: bytes, sheet_name: str = SHEET_MATERIAL) -> list[dict]:
    """原材料SC明細 シートをパースし、原料コード単位の行リストを返す。

    1原料 = 「数量」行 + 「金額」行。数量行から qty を、金額行から cost を読む。
    """
    wb = load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    ws = wb[sheet_name]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()

    by_code: dict[str, dict] = {}
    order: list[str] = []

    i = 4  # R5 (0-origin index 4)
    while i < len(rows):
        row = rows[i]
        if len(row) < 15:
            i += 1
            continue
        label = row[3]
        label_s = str(label).strip() if label is not None else ""
        if not label_s.startswith("数量"):
            i += 1
            continue
        # 数量行発見。コードを確認 (合計行は code が "合 計" などになる)
        code = _normalize_code(row[0])
        name = row[1]
        name_s = str(name).strip() if name else None
        if not code or (name_s is None and code in ("合計", "合　　　計")):
            i += 1
            continue
        if "合" in code and "計" in code:
            i += 1
            continue

        qty_row = row
        amount_row = rows[i + 1] if i + 1 < len(rows) else ()
        # 金額行か確認 (なければ amounts 全0扱い)
        amt_label = amount_row[3] if len(amount_row) > 3 else None
        if amt_label is None or not str(amt_label).strip().startswith("金額"):
            amount_row = ()  # 金額行が無い → cost 0
            advance = 1
        else:
            advance = 2

        adjustment_items: dict[str, dict[str, str]] = {}
        for adj_label, json_key, idx in ADJUSTMENT_ITEM_COLS:
            qv = _to_decimal(qty_row[idx]) if len(qty_row) > idx else Decimal("0")
            cv = _to_decimal(amount_row[idx]) if len(amount_row) > idx else Decimal("0")
            adjustment_items[json_key] = {
                "label": adj_label,
                "qty": str(qv),
                "cost": str(cv),
            }

        parsed = {
            "material_code": code,
            "material_name": name_s,
            "sc_unit_price": _to_decimal(row[2]),
            "unit": label_s or None,
            "adjustment_items": adjustment_items,
        }
        for prefix, idx in MAIN_GROUPS:
            parsed[f"{prefix}_qty"] = _to_decimal(qty_row[idx]) if len(qty_row) > idx else Decimal("0")
            parsed[f"{prefix}_cost"] = _to_decimal(amount_row[idx]) if len(amount_row) > idx else Decimal("0")

        if code in by_code:
            _merge_row(by_code[code], parsed)
        else:
            by_code[code] = parsed
            order.append(code)

        i += advance

    return [by_code[c] for c in order]


async def resolve_material_ids(
    db: AsyncSession, codes: set[str]
) -> dict[str, uuid.UUID]:
    """原料コードから material_id を解決。マスタにない場合は除外。"""
    if not codes:
        return {}
    res = await db.execute(
        select(Material.id, Material.code).where(Material.code.in_(codes))
    )
    return {code: mid for mid, code in res.all()}


async def process_prior_year_material_import(
    db: AsyncSession,
    file_content: bytes,
    filename: str,
    fiscal_year: int,
    sheet_name: str = SHEET_MATERIAL,
    source_system: str = "manual",
    delete_existing: bool = True,
    period_id: uuid.UUID | None = None,
) -> ImportBatch:
    """前年実績 原材料 Excel を取り込み、ImportBatch を返す。

    fiscal_year=38 で 10 SC原材料.xlsx を取り込む想定。
    delete_existing=True (デフォルト) で既存の同 fiscal_year レコードを
    全削除してから登録（マスタ的データなので冪等取込）。
    """
    batch = ImportBatch(
        file_name=filename,
        source_system=source_system if source_system in SourceSystem._value2member_map_
        else SourceSystem.manual.value,
        status=ImportStatus.processing,
        period_id=period_id,
        total_rows=0,
        success_rows=0,
        error_rows=0,
        started_at=datetime.now(),
    )
    db.add(batch)
    await db.flush()

    try:
        rows = parse_material_sheet(file_content, sheet_name)
    except Exception as e:
        batch.status = ImportStatus.failed
        batch.completed_at = datetime.now()
        batch.notes = f"ファイルパースエラー: {e}"
        db.add(ImportErrorModel(batch_id=batch.id, row_number=0, error_message=batch.notes))
        await db.flush()
        await db.refresh(batch)
        return batch

    batch.total_rows = len(rows)

    if delete_existing:
        await db.execute(
            delete(PriorYearMaterialActual).where(
                PriorYearMaterialActual.fiscal_year == fiscal_year
            )
        )
        await db.flush()

    codes = {r["material_code"] for r in rows}
    code_to_mid = await resolve_material_ids(db, codes)

    success = 0
    matched = 0
    unmatched = 0
    purchase_sum = Decimal("0")
    input_sum = Decimal("0")
    close_sum = Decimal("0")

    for r in rows:
        mid = code_to_mid.get(r["material_code"])
        if mid is not None:
            matched += 1
        else:
            unmatched += 1
        rec = PriorYearMaterialActual(
            fiscal_year=fiscal_year,
            material_code=r["material_code"],
            material_name=r["material_name"],
            material_id=mid,
            sc_unit_price=r["sc_unit_price"],
            unit=r["unit"],
            opening_qty=r["opening_qty"],
            opening_cost=r["opening_cost"],
            purchase_qty=r["purchase_qty"],
            purchase_cost=r["purchase_cost"],
            input_qty=r["input_qty"],
            input_cost=r["input_cost"],
            closing_qty=r["closing_qty"],
            closing_cost=r["closing_cost"],
            adjustment_items=r["adjustment_items"],
            source_file=filename,
            source_sheet=sheet_name,
            import_batch_id=batch.id,
        )
        db.add(rec)
        success += 1
        purchase_sum += r["purchase_cost"]
        input_sum += r["input_cost"]
        close_sum += r["closing_cost"]

    await db.flush()

    batch.success_rows = success
    batch.error_rows = 0
    batch.status = ImportStatus.completed
    batch.completed_at = datetime.now()
    batch.notes = (
        f"fiscal_year={fiscal_year}, total={len(rows)}, success={success}, "
        f"matched={matched}, unmatched={unmatched}; "
        f"purchase_sum={purchase_sum:.0f}, input_sum={input_sum:.0f}, "
        f"closing_sum={close_sum:.0f}"
    )
    await db.flush()
    await db.refresh(batch)
    return batch
