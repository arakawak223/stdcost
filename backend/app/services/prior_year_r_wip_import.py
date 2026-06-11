"""前年実績(F-01 横展開 最終) R仕掛品 加重平均原価コンポーネント取込サービス。

R仕掛品(原液系列 R1/R2/R3)の加重平均単価を導出するクロス集計ワークシートから、
各系列の **加重平均結果のみ** を prior_year_r_wip_components テーブルに upsert する。

対象2ファイル (cost_component で区別):
  - material: 21-1 R仕掛品　原材料.xlsx 「R仕掛原材料費」
      各系列ブロックの「加重平均」ラベル行 + 次2行 (金額/単価) を採用。
  - labor   : 21-2 R仕掛品　労務費.xlsx 「R仕掛品　労務費」
      系列ごとに複数の候補加重平均があり、注記「□を採用する」のとおり
      罫線で囲まれた採用ブロック (平均単価ラベルセルが left+bottom 罫線) を検出。

いずれも 3系列 (R1/R2/R3) × 数量(KG)/金額(円)/単価(円/KG)。fiscal_year 単位 (38期)。
"""

import io
import uuid
from datetime import datetime
from decimal import Decimal

from openpyxl import load_workbook
from sqlalchemy import delete
from sqlalchemy.ext.asyncio import AsyncSession

from app.models.audit import ImportBatch, ImportError as ImportErrorModel, ImportStatus
from app.models.cost import PriorYearRWipComponent, SourceSystem


SHEET_MATERIAL = "R仕掛原材料費"
SHEET_LABOR = "R仕掛品　労務費"

# 系列ブロックの (label列, value列) — openpyxl 1-origin
#   原材料: 加重平均ラベルは B/G/L 列、値は E/J/O 列
#   労務費: 平均単価ラベルは D/I/N 列、値は E/J/O 列
SERIES_BLOCKS_MATERIAL: list[tuple[str, int, int]] = [
    ("R1", 2, 5),
    ("R2", 7, 10),
    ("R3", 12, 15),
]
SERIES_BLOCKS_LABOR: list[tuple[str, int, int]] = [
    ("R1", 4, 5),
    ("R2", 9, 10),
    ("R3", 14, 15),
]


def _to_decimal(v) -> Decimal:
    if v is None:
        return Decimal("0")
    if isinstance(v, str):
        s = v.strip()
        if not s:
            return Decimal("0")
        try:
            return Decimal(s)
        except Exception:
            return Decimal("0")
    try:
        return Decimal(str(v))
    except Exception:
        return Decimal("0")


def parse_material_sheet(content: bytes, sheet_name: str = SHEET_MATERIAL) -> list[dict]:
    """R仕掛原材料費 シートから各系列の「加重平均」結果を抽出する。

    系列の先頭列に「加重平均」ラベルがある行を探し、その行=数量(KG)、+1行=金額、
    +2行=単価 を value列から読む。
    """
    wb = load_workbook(io.BytesIO(content), data_only=True)
    ws = wb[sheet_name]
    max_row = ws.max_row

    out: list[dict] = []
    for series, lcol, vcol in SERIES_BLOCKS_MATERIAL:
        anchor_row = None
        for r in range(1, max_row + 1):
            cell = ws.cell(r, lcol).value
            if cell is not None and str(cell).strip() == "加重平均":
                anchor_row = r
                break
        if anchor_row is None:
            continue
        qty = _to_decimal(ws.cell(anchor_row, vcol).value)
        amount = _to_decimal(ws.cell(anchor_row + 1, vcol).value)
        unit = _to_decimal(ws.cell(anchor_row + 2, vcol).value)
        out.append(
            {
                "r_series": series,
                "cost_component": "material",
                "weighted_avg_qty": qty,
                "weighted_avg_amount": amount,
                "weighted_avg_unit_price": unit,
            }
        )
    wb.close()
    return out


def parse_labor_sheet(content: bytes, sheet_name: str = SHEET_LABOR) -> list[dict]:
    """R仕掛品　労務費 シートから各系列の採用加重平均(罫線囲み)を抽出する。

    系列の label列に「平均単価」を含み、かつ left+bottom 罫線が引かれたセルが
    採用ブロックの最下行(単価)。その値=単価、-1行=金額、-2行=数量(KG)。
    """
    wb = load_workbook(io.BytesIO(content), data_only=True)
    ws = wb[sheet_name]
    max_row = ws.max_row

    out: list[dict] = []
    for series, lcol, vcol in SERIES_BLOCKS_LABOR:
        adopted_row = None
        for r in range(1, max_row + 1):
            cell = ws.cell(r, lcol)
            v = cell.value
            if (
                v is not None
                and "平均単価" in str(v)
                and cell.border.bottom.style
                and cell.border.left.style
            ):
                adopted_row = r  # 最後に見つかった囲みを採用(末尾の最終調整値)
        if adopted_row is None:
            continue
        unit = _to_decimal(ws.cell(adopted_row, vcol).value)
        amount = _to_decimal(ws.cell(adopted_row - 1, vcol).value)
        qty = _to_decimal(ws.cell(adopted_row - 2, vcol).value)
        out.append(
            {
                "r_series": series,
                "cost_component": "labor",
                "weighted_avg_qty": qty,
                "weighted_avg_amount": amount,
                "weighted_avg_unit_price": unit,
            }
        )
    wb.close()
    return out


async def process_prior_year_r_wip_import(
    db: AsyncSession,
    file_content: bytes,
    filename: str,
    fiscal_year: int,
    cost_component: str,
    sheet_name: str | None = None,
    source_system: str = "manual",
    delete_existing: bool = True,
    period_id: uuid.UUID | None = None,
) -> ImportBatch:
    """R仕掛品 加重平均コンポーネント Excel を取り込み、ImportBatch を返す。

    cost_component: 'material' (21-1) または 'labor' (21-2)。
    """
    if cost_component not in ("material", "labor"):
        raise ValueError(f"cost_component は material/labor のいずれか: {cost_component}")

    sheet = sheet_name or (SHEET_MATERIAL if cost_component == "material" else SHEET_LABOR)

    batch = ImportBatch(
        file_name=filename,
        source_system=source_system if source_system in SourceSystem._value2member_map_
        else SourceSystem.manual.value,
        status=ImportStatus.processing,
        period_id=period_id,
        total_rows=0,
        success_rows=0,
        error_rows=0,
        started_at=datetime.now(),
    )
    db.add(batch)
    await db.flush()

    try:
        if cost_component == "material":
            rows = parse_material_sheet(file_content, sheet)
        else:
            rows = parse_labor_sheet(file_content, sheet)
    except Exception as e:
        batch.status = ImportStatus.failed
        batch.completed_at = datetime.now()
        batch.notes = f"ファイルパースエラー: {e}"
        db.add(ImportErrorModel(batch_id=batch.id, row_number=0, error_message=batch.notes))
        await db.flush()
        await db.refresh(batch)
        return batch

    batch.total_rows = len(rows)

    if delete_existing:
        await db.execute(
            delete(PriorYearRWipComponent).where(
                PriorYearRWipComponent.fiscal_year == fiscal_year,
                PriorYearRWipComponent.cost_component == cost_component,
            )
        )
        await db.flush()

    success = 0
    amount_sum = Decimal("0")
    for r in rows:
        rec = PriorYearRWipComponent(
            fiscal_year=fiscal_year,
            r_series=r["r_series"],
            cost_component=r["cost_component"],
            weighted_avg_qty=r["weighted_avg_qty"],
            weighted_avg_amount=r["weighted_avg_amount"],
            weighted_avg_unit_price=r["weighted_avg_unit_price"],
            source_file=filename,
            source_sheet=sheet,
            import_batch_id=batch.id,
        )
        db.add(rec)
        success += 1
        amount_sum += r["weighted_avg_amount"]

    await db.flush()

    batch.success_rows = success
    batch.error_rows = 0
    batch.status = ImportStatus.completed
    batch.completed_at = datetime.now()
    batch.notes = (
        f"fiscal_year={fiscal_year}, component={cost_component}, "
        f"series={[r['r_series'] for r in rows]}, amount_sum={amount_sum:.0f}"
    )
    await db.flush()
    await db.refresh(batch)
    return batch
