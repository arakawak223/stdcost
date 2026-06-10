"""前年実績(F-01 横展開) 外注製品インポートサービス —
31SC外注製品.xlsx 「標準原価_外注製品」シートから 第38期の外注製品別
単価(38期実際/39期標準) + 数量・金額フローを prior_year_outsource_actuals に upsert。

シート構造 (標準原価_外注製品):
  R1-R4 マルチヘッダ
  R5    列ラベル
  R6〜  データ行 (1製品=1行)。区分見出し行 (c2 のみ "A.〜F."、c3-c5 空) で
        セクションが切り替わる (以降の製品行へ繰り下げ適用)。
  末尾 空行。

列マップ (1-origin):
  2: 区分見出し (A.外注先にて原料製造 等) ※見出し行のみ
  3: 外注先補助科目コード
  4: 外注先名
  5: 商品コード (=product_code, 一意キー)
  6: 外注製品名
  7: 支給原料の商品コード
  8: 支給原料 (ラベル)
  単価情報 単位原価/個:
    9 / 10 / 11  = 38期実際 外注加工費 / その他原価 / 外注製品原価
    13 / 14 / 15 = 39期標準 外注加工費 / その他原価 / 外注製品原価
  数量情報 (個): 17:37期末 18:生産 19:販売 20:外注支給 21:製品内製へ
    22:その他振替 23:販促DM 24:販促 25:試験研究費 26:交際費 27:寄付金
    28:広告宣伝費 29:在庫調整 30:38期末
  金額情報 (円): 31:37期末 32:生産 33:販売 34:外注支給 35:製品内製へ
    36:その他振替 37:販促DM 38:販促 39:試験研究費 40:交際費 41:寄付金
    42:広告宣伝費 43:在庫調整 44:38期末
"""

import io
import uuid
from datetime import datetime
from decimal import Decimal

from openpyxl import load_workbook
from sqlalchemy import delete, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.models.audit import ImportBatch, ImportError as ImportErrorModel, ImportStatus
from app.models.cost import PriorYearOutsourceActual, SourceSystem
from app.models.master import Product


SHEET_OUTSOURCE = "標準原価_外注製品"

# 主要4項目: (prefix, 数量列index(0-origin), 金額列index(0-origin))
MAIN_GROUPS: list[tuple[str, int, int]] = [
    ("opening", 16, 30),     # 37期末
    ("production", 17, 31),  # 生産
    ("sales", 18, 32),       # 販売
    ("closing", 29, 43),     # 38期末
]

# その他10項目: (項目名, JSONキー, 数量列index, 金額列index)
MOVEMENT_COLS: list[tuple[str, str, int, int]] = [
    ("外注支給", "outsource_supply", 19, 33),
    ("製品内製へ", "to_internal", 20, 34),
    ("その他振替", "other_transfer", 21, 35),
    ("販促DM", "promo_dm", 22, 36),
    ("販促", "promo", 23, 37),
    ("試験研究費", "rnd", 24, 38),
    ("交際費", "entertainment", 25, 39),
    ("寄付金", "donation", 26, 40),
    ("広告宣伝費", "advertising", 27, 41),
    ("在庫調整", "inventory_adj", 28, 42),
]


def _to_decimal(v) -> Decimal:
    if v is None:
        return Decimal("0")
    if isinstance(v, bool):
        return Decimal("0")
    if isinstance(v, str):
        s = v.strip()
        if not s or s in ("ー", "−", "-", "#REF!", "#N/A", "#VALUE!"):
            return Decimal("0")
        try:
            return Decimal(s)
        except Exception:
            return Decimal("0")
    try:
        return Decimal(str(v))
    except Exception:
        return Decimal("0")


def _str_or_none(v) -> str | None:
    if v is None:
        return None
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    s = str(v).strip()
    return s if s else None


def _cell(row, idx) -> Decimal:
    return _to_decimal(row[idx]) if len(row) > idx else Decimal("0")


def parse_outsource_sheet(content: bytes, sheet_name: str = SHEET_OUTSOURCE) -> list[dict]:
    """標準原価_外注製品 シートをパースし、外注製品行リストを返す。

    区分見出し行 (c2 のみ) は current_section を更新してスキップ。
    商品コード(c5) を持つ行を製品行として取り込む。
    """
    wb = load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    ws = wb[sheet_name]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()

    out: list[dict] = []
    current_section: str | None = None

    for ri, row in enumerate(rows):
        if ri < 5:  # R1-R5 はヘッダ
            continue
        if len(row) < 16:
            continue
        section_label = _str_or_none(row[1])  # c2
        product_code = _str_or_none(row[4])   # c5
        if product_code is None:
            # 区分見出し行 (c2 のみ) ならセクション更新
            if section_label:
                current_section = section_label
            continue

        movements: dict[str, dict[str, str]] = {}
        for label, json_key, qidx, cidx in MOVEMENT_COLS:
            movements[json_key] = {
                "label": label,
                "qty": str(_cell(row, qidx)),
                "cost": str(_cell(row, cidx)),
            }

        parsed = {
            "product_code": product_code,
            "product_name": _str_or_none(row[5]),
            "section": current_section,
            "contractor_code": _str_or_none(row[2]),
            "contractor_name": _str_or_none(row[3]),
            "supplied_material_code": _str_or_none(row[6]),
            "supplied_material_label": _str_or_none(row[7]),
            "actual_processing_cost": _cell(row, 8),
            "actual_other_cost": _cell(row, 9),
            "actual_product_cost": _cell(row, 10),
            "std_processing_cost": _cell(row, 12),
            "std_other_cost": _cell(row, 13),
            "std_product_cost": _cell(row, 14),
            "movements": movements,
        }
        for prefix, qidx, cidx in MAIN_GROUPS:
            parsed[f"{prefix}_qty"] = _cell(row, qidx)
            parsed[f"{prefix}_cost"] = _cell(row, cidx)

        out.append(parsed)

    return out


async def resolve_product_ids(
    db: AsyncSession, codes: set[str]
) -> dict[str, uuid.UUID]:
    if not codes:
        return {}
    res = await db.execute(
        select(Product.id, Product.code).where(Product.code.in_(codes))
    )
    return {code: pid for pid, code in res.all()}


async def process_prior_year_outsource_import(
    db: AsyncSession,
    file_content: bytes,
    filename: str,
    fiscal_year: int,
    sheet_name: str = SHEET_OUTSOURCE,
    source_system: str = "manual",
    delete_existing: bool = True,
    period_id: uuid.UUID | None = None,
) -> ImportBatch:
    """前年実績 外注製品 Excel を取り込み、ImportBatch を返す。"""
    batch = ImportBatch(
        file_name=filename,
        source_system=source_system if source_system in SourceSystem._value2member_map_
        else SourceSystem.manual.value,
        status=ImportStatus.processing,
        period_id=period_id,
        total_rows=0,
        success_rows=0,
        error_rows=0,
        started_at=datetime.now(),
    )
    db.add(batch)
    await db.flush()

    try:
        rows = parse_outsource_sheet(file_content, sheet_name)
    except Exception as e:
        batch.status = ImportStatus.failed
        batch.completed_at = datetime.now()
        batch.notes = f"ファイルパースエラー: {e}"
        db.add(ImportErrorModel(batch_id=batch.id, row_number=0, error_message=batch.notes))
        await db.flush()
        await db.refresh(batch)
        return batch

    batch.total_rows = len(rows)

    if delete_existing:
        await db.execute(
            delete(PriorYearOutsourceActual).where(
                PriorYearOutsourceActual.fiscal_year == fiscal_year
            )
        )
        await db.flush()

    codes = {r["product_code"] for r in rows}
    code_to_pid = await resolve_product_ids(db, codes)

    success = 0
    matched = 0
    seen: set[str] = set()
    dup = 0
    by_section: dict[str, int] = {}
    prod_sum = Decimal("0")
    sales_sum = Decimal("0")
    close_sum = Decimal("0")

    for r in rows:
        if r["product_code"] in seen:
            dup += 1
            continue
        seen.add(r["product_code"])
        pid = code_to_pid.get(r["product_code"])
        if pid is not None:
            matched += 1
        rec = PriorYearOutsourceActual(
            fiscal_year=fiscal_year,
            product_code=r["product_code"],
            product_name=r["product_name"],
            product_id=pid,
            section=r["section"],
            contractor_code=r["contractor_code"],
            contractor_name=r["contractor_name"],
            supplied_material_code=r["supplied_material_code"],
            supplied_material_label=r["supplied_material_label"],
            actual_processing_cost=r["actual_processing_cost"],
            actual_other_cost=r["actual_other_cost"],
            actual_product_cost=r["actual_product_cost"],
            std_processing_cost=r["std_processing_cost"],
            std_other_cost=r["std_other_cost"],
            std_product_cost=r["std_product_cost"],
            opening_qty=r["opening_qty"],
            opening_cost=r["opening_cost"],
            production_qty=r["production_qty"],
            production_cost=r["production_cost"],
            sales_qty=r["sales_qty"],
            sales_cost=r["sales_cost"],
            closing_qty=r["closing_qty"],
            closing_cost=r["closing_cost"],
            movements=r["movements"],
            source_file=filename,
            source_sheet=sheet_name,
            import_batch_id=batch.id,
        )
        db.add(rec)
        success += 1
        sec = r["section"] or "(none)"
        by_section[sec] = by_section.get(sec, 0) + 1
        prod_sum += r["production_cost"]
        sales_sum += r["sales_cost"]
        close_sum += r["closing_cost"]

    await db.flush()

    batch.success_rows = success
    batch.error_rows = 0
    batch.status = ImportStatus.completed
    batch.completed_at = datetime.now()
    sec_summary = ", ".join(f"{k}={v}" for k, v in sorted(by_section.items()))
    batch.notes = (
        f"fiscal_year={fiscal_year}, total={len(rows)}, success={success}, "
        f"matched={matched}, dup_skipped={dup}; sections: {sec_summary}; "
        f"production_sum={prod_sum:.0f}, sales_sum={sales_sum:.0f}, "
        f"closing_sum={close_sum:.0f}"
    )
    await db.flush()
    await db.refresh(batch)
    return batch
